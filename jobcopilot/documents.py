"""ATS-friendly DOCX and PDF rendering.

Rules applied: single column, standard section headings, real text (no images,
tables or text boxes), standard fonts (Calibri in DOCX; in PDF a Unicode font such as
DejaVu Sans when installed so non-Latin names render, otherwise Helvetica),
simple bullets, and contact details in the body rather than a header/footer
(some ATS parsers skip headers).
"""

from __future__ import annotations

import os
from functools import cache
from pathlib import Path
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .profile import strip_control_chars

if TYPE_CHECKING:
    from .tailoring import TailoredResume

FONT = "Calibri"

_WINDOWS_FONTS = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
# (regular, bold) candidates for a TrueType font with wide Unicode coverage.
_UNICODE_FONTS = [
    (d / "DejaVuSans.ttf", d / "DejaVuSans-Bold.ttf")
    for d in (
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts/dejavu"),
        Path("/usr/share/fonts/TTF"),
        Path("/usr/local/share/fonts"),
        Path("/Library/Fonts"),
        Path.home() / "Library/Fonts",
        Path.home() / ".local/share/fonts",
        _WINDOWS_FONTS,
    )
] + [
    (Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"), None),
    (Path("/Library/Fonts/Arial Unicode.ttf"), None),
    (_WINDOWS_FONTS / "arialuni.ttf", None),
]


@cache
def pdf_fonts() -> tuple[str, str]:
    """(regular, bold) PDF font names: a Unicode TTF if one is installed, else Helvetica."""
    for regular, bold in _UNICODE_FONTS:
        if not regular.is_file():
            continue
        try:
            pdfmetrics.registerFont(TTFont("ResumeSans", str(regular)))
            bold_name = "ResumeSans"
            if bold is not None and bold.is_file():
                pdfmetrics.registerFont(TTFont("ResumeSans-Bold", str(bold)))
                bold_name = "ResumeSans-Bold"
            return "ResumeSans", bold_name
        except Exception:  # unreadable or unsupported font file: try the next one
            continue
    return "Helvetica", "Helvetica-Bold"


def _clean(text: str) -> str:
    """Strip characters DOCX/PDF cannot hold (python-docx raises on control chars)."""
    return strip_control_chars(text or "")


def _clean_resume(r: TailoredResume) -> TailoredResume:
    """A copy of the resume with every string sanitized for rendering."""

    def walk(value):
        if isinstance(value, str):
            return _clean(value)
        if isinstance(value, list):
            return [walk(v) for v in value]
        if isinstance(value, dict):
            return {k: walk(v) for k, v in value.items()}
        return value

    return type(r).model_validate(walk(r.model_dump()))


def _sections(r: TailoredResume) -> list[tuple[str, list[tuple[str, list[str]]]]]:
    """Resume as (heading, [(line, bullets)]) so DOCX and PDF share one layout."""
    out: list[tuple[str, list[tuple[str, list[str]]]]] = []
    if r.summary:
        out.append(("SUMMARY", [(r.summary, [])]))
    if r.skills:
        out.append(("SKILLS", [(", ".join(r.skills), [])]))
    if r.experience:
        out.append(
            (
                "EXPERIENCE",
                [
                    (f"{e.title} | {e.company} | {e.location} | {e.start} - {e.end}", e.bullets)
                    for e in r.experience
                ],
            )
        )
    if r.projects:
        out.append(("PROJECTS", [(p.name, p.bullets) for p in r.projects]))
    if r.education:
        out.append(
            (
                "EDUCATION",
                [(f"{e.credential} | {e.institution} | {e.dates}", e.details) for e in r.education],
            )
        )
    if r.certifications:
        out.append(("CERTIFICATIONS", [(c, []) for c in r.certifications]))
    return out


def _base_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)
    return doc


def render_resume_docx(r: TailoredResume, path: Path) -> Path:
    """Write the resume as a single-column DOCX."""
    r = _clean_resume(r)
    doc = _base_document()
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(r.name)
    run.bold = True
    run.font.size = Pt(18)
    for line in (r.headline, " | ".join(r.contact)):
        if line:
            p = doc.add_paragraph(line)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, entries in _sections(r):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        hr = h.add_run(heading)
        hr.bold = True
        hr.font.size = Pt(12)
        for line, bullets in entries:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = heading in {"EXPERIENCE", "PROJECTS", "EDUCATION"}
            for b in bullets:
                doc.add_paragraph(b, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def render_resume_pdf(r: TailoredResume, path: Path) -> Path:
    """Write the resume as a text-based (selectable, parseable) single-column PDF."""
    r = _clean_resume(r)
    regular_font, bold_font = pdf_fonts()
    body = ParagraphStyle("body", fontName=regular_font, fontSize=10.5, leading=13.5)
    bold = ParagraphStyle("bold", parent=body, fontName=bold_font)
    heading = ParagraphStyle("h", parent=bold, fontSize=12, spaceBefore=10, spaceAfter=3)
    name = ParagraphStyle("name", parent=bold, fontSize=18, leading=22, alignment=TA_CENTER)
    center = ParagraphStyle("center", parent=body, alignment=TA_CENTER)

    story: list = [Paragraph(escape(r.name), name)]
    for line in (r.headline, " | ".join(r.contact)):
        if line:
            story.append(Paragraph(escape(line), center))
    for title, entries in _sections(r):
        story.append(Paragraph(title, heading))
        for line, bullets in entries:
            style = bold if title in {"EXPERIENCE", "PROJECTS", "EDUCATION"} else body
            story.append(Paragraph(escape(line), style))
            if bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(escape(b), body), leftIndent=12) for b in bullets],
                        bulletType="bullet",
                        start="•",
                        leftIndent=12,
                    )
                )
            story.append(Spacer(1, 3))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=0.8 * inch,
        rightMargin=0.8 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=f"{r.name} - Resume",
        author=r.name,
    ).build(story)
    return path


def render_cover_letter_docx(
    greeting: str, paragraphs: list[str], closing: str, path: Path
) -> Path:
    """Write a plain cover letter DOCX."""
    doc = _base_document()
    for text in [greeting, *paragraphs, closing]:
        doc.add_paragraph(_clean(text))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
