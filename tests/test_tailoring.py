import io

from conftest import FakeLLM, make_job
from docx import Document
from pypdf import PdfReader

from jobcopilot.tailoring import (
    AnswerItem,
    ApplicationAnswers,
    CoverLetter,
    EducationEntry,
    ExperienceEntry,
    JobRequirements,
    TailoredResume,
    answer_questions,
    draft_cover_letter,
    keyword_coverage,
    normalize,
    tailor_resume,
    verify_resume,
)

REQS = JobRequirements(
    required_skills=["Python", "FastAPI", "Kubernetes"],
    nice_to_have=["Go"],
    keywords=["REST APIs", "PostgreSQL", "CI/CD"],
)


def tailored(**overrides) -> TailoredResume:
    data = dict(
        name="Alex Rivera",
        headline="Backend Python Developer",
        contact=["alex.rivera@example.com", "Toronto, ON"],
        summary="Backend developer with 4 years of Python, FastAPI and PostgreSQL experience.",
        skills=[
            "Python",
            "FastAPI",
            "PostgreSQL",
            "REST APIs",
            "Kubernetes",
        ],  # Kubernetes: invented
        experience=[
            ExperienceEntry(
                title="Backend Developer",
                company="ShipFast Logistics",
                location="Toronto, ON",
                start="Mar 2022",
                end="Present",
                bullets=[
                    "Built a FastAPI quote service handling about 40,000 requests per day.",
                    "Cut infrastructure costs by 35%.",
                ],  # 35% is not in the profile
            ),
            ExperienceEntry(
                title="Engineer",
                company="Initech",
                location="Remote",
                start="2019",
                end="2020",
                bullets=["Invented employer"],
            ),
        ],
        projects=[],
        education=[
            EducationEntry(
                credential="B.Sc. Computer Science",
                institution="University of Waterloo",
                dates="2016 - 2020",
                details=[],
            )
        ],
        certifications=["AWS Certified Cloud Practitioner", "CKA"],
        review_items=["Kubernetes experience is not in the profile"],
        change_log=["Moved FastAPI work to the top"],
    )
    data.update(overrides)
    return TailoredResume(**data)


def test_normalize_treats_punctuation_consistently():
    assert normalize("Node.js") == normalize("node js")
    assert "c++" in normalize("C++ and Rust")


def test_verify_resume_removes_fabrications(profile):
    resume = tailored()
    report = verify_resume(resume, profile)
    assert [e.company for e in resume.experience] == ["ShipFast Logistics"]
    assert "Kubernetes" not in resume.skills
    assert resume.certifications == ["AWS Certified Cloud Practitioner"]
    assert "Employer not in profile: Initech" in report.removed
    assert "Skill not in profile: Kubernetes" in report.removed
    assert "Certification not in profile: CKA" in report.removed
    assert any("35%" in w for w in report.warnings)
    assert not any("40,000" in w for w in report.warnings)  # real metric passes


def test_keyword_coverage():
    cov = keyword_coverage(["Python", "Kubernetes", "python"], "I write Python daily")
    assert cov == {"percent": 50, "matched": ["Python"], "missing": ["Kubernetes"]}


def test_tailor_resume_pipeline(db, profile, tmp_path):
    job = make_job(id="job1")
    db.upsert_jobs([job])
    llm = FakeLLM({JobRequirements: REQS, TailoredResume: tailored()})
    result = tailor_resume(db, llm, profile, job, tmp_path / "out")

    # Two Claude calls: requirement extraction, then the rewrite with target keywords.
    assert [c["schema"] for c in llm.calls] == [JobRequirements, TailoredResume]
    assert "Kubernetes" in llm.calls[1]["prompt"]

    kw = result["keywords"]
    assert "Kubernetes" in kw["not_in_profile"]
    assert "Kubernetes" not in kw["after"]["matched"]  # never added, it was stripped
    assert kw["after"]["percent"] >= kw["before"]["percent"]
    assert result["verification"]["removed"]

    out = tmp_path / "out" / "job1"
    docx_text = "\n".join(
        p.text
        for p in Document(io.BytesIO((out / result["files"]["docx"]).read_bytes())).paragraphs
    )
    assert "EXPERIENCE" in docx_text and "ShipFast Logistics" in docx_text
    assert "Initech" not in docx_text
    pdf_text = "".join(p.extract_text() for p in PdfReader(out / result["files"]["pdf"]).pages)
    assert "Alex Rivera" in pdf_text and "SKILLS" in pdf_text

    stored = db.latest_artifact("job1", "resume")
    assert stored["files"] == result["files"]


def test_ats_docx_has_no_tables_or_images(db, profile, tmp_path):
    job = make_job(id="job2")
    db.upsert_jobs([job])
    llm = FakeLLM({JobRequirements: REQS, TailoredResume: tailored()})
    result = tailor_resume(db, llm, profile, job, tmp_path)
    doc = Document(str(tmp_path / "job2" / result["files"]["docx"]))
    assert doc.tables == []
    assert doc.inline_shapes._inline_lst == []
    assert len(doc.sections) == 1


def test_cover_letter_is_flagged_for_review(db, profile, tmp_path):
    job = make_job(id="job3")
    db.upsert_jobs([job])
    letter = CoverLetter(
        greeting="Dear Hiring Team,",
        paragraphs=["I build Python APIs."],
        closing="Sincerely,\nAlex Rivera",
        review_items=["Confirm salary"],
    )
    result = draft_cover_letter(db, FakeLLM({CoverLetter: letter}), profile, job, tmp_path)
    assert result["review_required"] is True
    assert "I build Python APIs." in result["text"]
    assert (tmp_path / "job3" / result["files"]["docx"]).exists()


def test_unanswered_questions_are_flagged(db, profile):
    job = make_job(id="job4")
    db.upsert_jobs([job])
    answers = ApplicationAnswers(
        answers=[
            AnswerItem(
                question="Notice period?",
                answer="Two weeks.",
                source="EXPERIENCE_ANSWERS.md",
                needs_review=False,
            ),
            AnswerItem(question="Expected salary?", answer="", source="", needs_review=False),
        ]
    )
    llm = FakeLLM({ApplicationAnswers: answers})
    result = answer_questions(db, llm, profile, job, ["Notice period?", "Expected salary?", ""])
    assert [a["needs_review"] for a in result["answers"]] == [False, True]
    assert "- Expected salary?" in llm.calls[0]["prompt"]
