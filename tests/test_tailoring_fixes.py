"""Regression tests for the no-fabrication verifier, rendering and upload fixes."""

import io
from pathlib import Path

import pytest
from conftest import FakeLLM, make_job
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from jobcopilot import documents
from jobcopilot.profile import (
    NO_TEXT_MESSAGE,
    Profile,
    ProfileError,
    extract_text,
    load_profile,
    save_resume_upload,
)
from jobcopilot.tailoring import (
    AnswerItem,
    ApplicationAnswers,
    CoverLetter,
    EducationEntry,
    ExperienceEntry,
    ProjectEntry,
    TailoredResume,
    answer_questions,
    contains,
    draft_cover_letter,
    keyword_coverage,
    normalize,
    verify_resume,
)


def resume(**kw) -> TailoredResume:
    base = dict(
        name="Alex Rivera",
        headline="Backend Python Developer",
        contact=["alex.rivera@example.com"],
        summary="",
        skills=[],
        experience=[],
        projects=[],
        education=[],
        certifications=[],
        review_items=[],
        change_log=[],
    )
    base.update(kw)
    return TailoredResume(**base)


def job_entry(company="ShipFast Logistics", **kw) -> ExperienceEntry:
    data = dict(
        title="Backend Developer",
        company=company,
        location="Toronto, ON",
        start="Mar 2022",
        end="Present",
        bullets=[],
    )
    data.update(kw)
    return ExperienceEntry(**data)


def make_profile(resume_text: str, answers: str = "", criteria: str = "") -> Profile:
    return Profile(
        resume=resume_text,
        criteria=criteria,
        answers=answers,
        directory=Path("."),
        resume_file="resume.md",
    )


# 1. numbers -----------------------------------------------------------------------


def test_number_check_uses_whole_numbers_not_substrings(profile):
    r = resume(
        summary="Cut latency by 20% for 4 years.",  # '20' only exists inside '2020'
        experience=[job_entry(bullets=["Handled about 40,000 requests per day for 12 carriers."])],
    )
    warnings = verify_resume(r, profile).warnings
    assert any("'20%'" in w for w in warnings)
    assert not any("40,000" in w or "'12'" in w or "'4'" in w for w in warnings)


# 2. employers / schools / certs / skills ---------------------------------------------


def test_skills_from_criteria_are_not_facts(profile):
    r = resume(skills=["Python", "security clearance", "Pure front-end", "Distributed Systems"])
    report = verify_resume(r, profile)
    assert r.skills == ["Python", "Distributed Systems"]
    assert "Skill not in profile: security clearance" in report.removed


def test_partial_names_are_flagged_not_deleted(profile):
    r = resume(
        experience=[job_entry(company="Logistics")],
        education=[
            EducationEntry(
                credential="B.Sc. Computer Science",
                institution="Waterloo",
                dates="2016 - 2020",
                details=[],
            )
        ],
        certifications=["AWS Certified"],
    )
    report = verify_resume(r, profile)
    assert [e.company for e in r.experience] == ["Logistics"]
    assert any("incomplete" in w and "Logistics" in w for w in report.warnings)
    assert any("incomplete" in w and "Waterloo" in w for w in report.warnings)
    assert any("incomplete" in w and "AWS Certified" in w for w in report.warnings)


def test_full_names_pass_cleanly(profile):
    r = resume(
        experience=[job_entry(), job_entry("Maple Retail Co.", start="Jun 2020", end="Feb 2022")],
        certifications=["AWS Certified Cloud Practitioner"],
    )
    report = verify_resume(r, profile)
    assert report.removed == [] and report.warnings == []


def test_employer_only_in_criteria_is_removed():
    p = make_profile("Worked at Acme Corp.", criteria="Dream employer: Globex")
    r = resume(experience=[job_entry("Globex", title="Worked")])
    report = verify_resume(r, p)
    assert r.experience == []
    assert "Employer not in profile: Globex" in report.removed


# 3. titles, dates, credentials, projects, headline -------------------------------------


def test_unchecked_fields_are_now_flagged(profile):
    r = resume(
        headline="Engineer with 15 years of experience",
        experience=[job_entry(title="Principal Engineer", start="2010")],
        education=[
            EducationEntry(
                credential="PhD Machine Learning",
                institution="University of Waterloo",
                dates="2016 - 2020",
                details=[],
            )
        ],
        projects=[ProjectEntry(name="Kubernetes Operator", bullets=["Served 9 clients"])],
    )
    w = verify_resume(r, profile).warnings
    assert "Title not in profile: Principal Engineer" in w
    assert "Date not in profile: 2010" in w
    assert "Credential not in profile: PhD Machine Learning" in w
    assert "Project not in profile: Kubernetes Operator" in w
    assert any("'15'" in x for x in w) and any("'9'" in x for x in w)
    assert len(r.experience) == 1 and len(r.education) == 1  # flagged, not deleted


# 4. unicode normalize ---------------------------------------------------------------------


def test_normalize_is_unicode_aware():
    assert normalize("株式会社サンプル").strip() != ""
    assert normalize("Université de Montréal") == normalize("universite de montreal")
    assert normalize("Johnson & Johnson") == normalize("Johnson and Johnson")
    assert contains(" foo ", "—", "a — b")  # empty needle falls back to raw text


def test_cjk_and_accented_names_survive_verification():
    p = make_profile("Worked at 株式会社サンプル, Université de Montréal and Johnson & Johnson.")
    r = resume(
        experience=[
            job_entry("株式会社サンプル", title="Worked"),
            job_entry("Universite de Montreal", title="Worked"),
            job_entry("Johnson and Johnson", title="Worked"),
        ]
    )
    report = verify_resume(r, p)
    assert len(r.experience) == 3 and report.removed == []
    assert not any("incomplete" in w for w in report.warnings)


# 5/6. rendering ---------------------------------------------------------------------------


def test_control_characters_do_not_crash_rendering(tmp_path):
    r = resume(summary="Built pipelines\x0cfor data\x00 and\x0bmore", skills=["Py\x07thon"])
    documents.render_resume_docx(r, tmp_path / "r.docx")
    documents.render_resume_pdf(r, tmp_path / "r.pdf")
    documents.render_cover_letter_docx("Hi\x0c", ["p\x00"], "c", tmp_path / "c.docx")
    text = "\n".join(p.text for p in Document(str(tmp_path / "r.docx")).paragraphs)
    assert "Python" in text and "\x00" not in text


def test_pdf_uses_unicode_font_when_available(tmp_path):
    regular, _ = documents.pdf_fonts()
    if regular == "Helvetica":
        pytest.skip("no Unicode TTF installed on this machine")
    documents.render_resume_pdf(resume(name="Zoë Ωmega Иван"), tmp_path / "r.pdf")
    text = PdfReader(tmp_path / "r.pdf").pages[0].extract_text()
    assert "Ωmega" in text and "Иван" in text


# 8. uploads ---------------------------------------------------------------------------------


def _blank_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(10, 10, 100, 100)
    c.showPage()
    c.save()
    return buf.getvalue()


GOOD = b"Jane Doe\nPython developer at Example Corp, 2020 - Present. Built APIs and pipelines."


@pytest.mark.parametrize(
    ("name", "data"),
    [("scan.pdf", _blank_pdf()), ("r.txt", b""), ("r.md", b"   \n  short  ")],
)
def test_upload_without_text_is_rejected_and_keeps_old_resume(tmp_path, name, data):
    save_resume_upload(tmp_path, "resume.md", GOOD)
    with pytest.raises(ProfileError, match="Scanned PDFs"):
        save_resume_upload(tmp_path, name, data)
    assert NO_TEXT_MESSAGE.startswith("Couldn't find text")
    assert load_profile(tmp_path).resume == GOOD.decode()


def test_binary_named_markdown_is_rejected(tmp_path):
    save_resume_upload(tmp_path, "resume.md", GOOD)
    with pytest.raises(ProfileError):
        save_resume_upload(tmp_path, "r.md", b"\x89PNG\r\n\x1a\n" + b"\x00\xff" * 100)
    with pytest.raises(ProfileError):
        save_resume_upload(tmp_path, "RESUME.PDF", b"not a pdf at all")
    assert sorted(p.name for p in tmp_path.iterdir()) == ["resume.md"]


def test_example_profile_flag_is_accurate(tmp_path):
    assert load_profile(tmp_path / "missing").is_example
    assert load_profile(Path("profile.example")).is_example
    save_resume_upload(tmp_path, "cv.txt", GOOD)
    assert not load_profile(tmp_path).is_example


# 9. DOCX tables, headers, footers -------------------------------------------------------------


def test_docx_extract_includes_tables_headers_and_footers():
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "ShipFast Logistics"
    t.cell(0, 1).text = "Python, FastAPI"
    merged = t.cell(1, 0).merge(t.cell(1, 1))
    merged.text = "Merged cell"
    nested = t.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested skill"
    doc.sections[0].header.paragraphs[0].text = "alex@example.com"
    doc.sections[0].footer.paragraphs[0].text = "Page footer"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("r.docx", buf.getvalue())
    for needle in ("alex@example.com", "ShipFast Logistics", "Nested skill", "Page footer"):
        assert needle in text
    assert text.count("Merged cell") == 1


# 10. keyword coverage -------------------------------------------------------------------------


def test_short_keyword_false_positives():
    text = "Led R&D; go-to-market; Series C; net revenue; JavaScript"
    cov = keyword_coverage(["R", "Go", "C", ".NET", "Java"], text)
    assert cov["matched"] == []


def test_short_keywords_still_match_real_usage():
    text = "Languages: C, R and Go. Built .NET services with C# and Node.js"
    cov = keyword_coverage(["C", "R", "Go", ".NET", "C#", "NodeJS", "node js"], text)
    assert cov["missing"] == []


# 11. answers and cover letter ----------------------------------------------------------------


def test_blank_questions_make_no_model_call(db, profile):
    db.upsert_jobs([make_job(id="q0")])
    llm = FakeLLM({})
    assert answer_questions(db, llm, profile, make_job(id="q0"), ["", "  "])["answers"] == []
    assert llm.calls == []


def test_answers_remapped_sensitive_and_numbers_flagged(db, profile):
    db.upsert_jobs([make_job(id="q1")])
    answers = ApplicationAnswers(
        answers=[
            AnswerItem(
                question="What is your notice period?",
                answer="Two weeks.",
                source="EXPERIENCE_ANSWERS.md",
                needs_review=False,
            ),
            AnswerItem(
                question="Do you need visa sponsorship?",
                answer="No.",
                source="EXPERIENCE_ANSWERS.md",
                needs_review=False,
            ),
            AnswerItem(
                question="Years of Python?",
                answer="About 9 years.",
                source="EXPERIENCE_ANSWERS.md",
                needs_review=False,
            ),
        ]
    )
    qs = [
        "  What is your notice period?",
        "Do you need visa sponsorship?",
        "Years of Python?",
        "Why us?",
    ]
    result = answer_questions(
        db, FakeLLM({ApplicationAnswers: answers}), profile, make_job(id="q1"), qs
    )
    got = result["answers"]
    assert [a["question"] for a in got] == [q.strip() for q in qs]
    assert [a["needs_review"] for a in got] == [False, True, True, True]
    assert got[3]["answer"] == ""
    assert any("'9'" in w for w in result["warnings"])


def test_cover_letter_numbers_are_flagged(db, profile, tmp_path):
    db.upsert_jobs([make_job(id="cl")])
    letter = CoverLetter(
        greeting="Dear team,",
        paragraphs=["I cut costs by 35% and mentored two developers over 4 years."],
        closing="Alex",
        review_items=[],
    )
    result = draft_cover_letter(
        db, FakeLLM({CoverLetter: letter}), profile, make_job(id="cl"), tmp_path
    )
    assert any("35%" in w for w in result["warnings"])
    assert any("35%" in w for w in result["letter"]["review_items"])
    assert not any("'4'" in w for w in result["warnings"])
